import os
import qrcode
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Image as RLImage, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

STORE_URL = "https://www.t3store.store"
QR_PATH = "t3_store_qr.png"
PDF_PATH = "T3_Superstore_Printable_Flyer.pdf"
DOCX_PATH = "T3_Superstore_Printable_Flyer.docx"
HTML_PATH = "T3_Superstore_Printable_Flyer.html"

# 1. Generate High-Resolution QR Code
def generate_qr():
    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.constants.ERROR_CORRECT_H,
        box_size=16,
        border=3,
    )
    qr.add_data(STORE_URL)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#0C1E3C", back_color="white").convert('RGB')
    img.save(QR_PATH, "PNG", dpi=(300, 300))
    print(f"Generated QR Code: {QR_PATH}")

# 2. Generate Professional ReportLab PDF Flyer (A4)
def generate_pdf():
    # A4 is 595.27 x 841.89 points
    c = canvas.Canvas(PDF_PATH, pagesize=A4)
    width, height = A4

    # Background
    c.setFillColor(colors.HexColor("#F8FAFC"))
    c.rect(0, 0, width, height, fill=1, stroke=0)

    # Outer decorative frame
    margin = 18
    c.setStrokeColor(colors.HexColor("#CBD5E1"))
    c.setLineWidth(1)
    c.roundRect(margin, margin, width - 2 * margin, height - 2 * margin, 14, fill=0, stroke=1)

    # Inner decorative border
    c.setStrokeColor(colors.HexColor("#E2E8F0"))
    c.setLineWidth(0.6)
    c.roundRect(margin + 4, margin + 4, width - 2 * (margin + 4), height - 2 * (margin + 4), 10, fill=0, stroke=1)

    # Header Navy Banner
    banner_y = height - margin - 110
    banner_height = 106
    c.setFillColor(colors.HexColor("#0C1E3C"))
    c.rect(margin + 5, banner_y, width - 2 * (margin + 5), banner_height, fill=1, stroke=0)

    # Top Red accent line on banner
    c.setFillColor(colors.HexColor("#E4002B"))
    c.rect(margin + 5, height - margin - 8, width - 2 * (margin + 5), 7, fill=1, stroke=0)

    # Header Brand Text
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 30)
    c.drawCentredString(width / 2, height - margin - 44, "T3 SUPERSTORE")

    c.setFillColor(colors.HexColor("#F59E0B"))  # Gold
    c.setFont("Helvetica-Bold", 12)
    c.drawCentredString(width / 2, height - margin - 64, "WHOLESALE DRINKS, PACKS & CRATES")

    c.setFillColor(colors.HexColor("#E2E8F0"))
    c.setFont("Helvetica", 9.5)
    c.drawCentredString(width / 2, height - margin - 86, "Direct Distributor Prices  ·  Event Supply  ·  Strictly Packs & Crates")

    # Announcement Pill Badge
    pill_w = 280
    pill_h = 26
    pill_x = (width - pill_w) / 2
    pill_y = banner_y - 28
    c.setFillColor(colors.HexColor("#E4002B"))
    c.roundRect(pill_x, pill_y, pill_w, pill_h, 13, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont("Helvetica-Bold", 11)
    c.drawCentredString(width / 2, pill_y + 8, "★ EXCITING NEWS FOR OUR CUSTOMERS ★")

    # Main Headline
    c.setFillColor(colors.HexColor("#0C1E3C"))
    c.setFont("Helvetica-Bold", 24)
    c.drawCentredString(width / 2, pill_y - 34, "YOU CAN NOW BOOK ONLINE!")

    # Sub-headline
    c.setFillColor(colors.HexColor("#475569"))
    c.setFont("Helvetica", 11.5)
    c.drawCentredString(width / 2, pill_y - 54, "Browse our complete catalog with live prices and order directly from your phone.")

    # Center QR Card Container
    qr_card_w = 340
    qr_card_h = 320
    qr_card_x = (width - qr_card_w) / 2
    qr_card_y = pill_y - 70 - qr_card_h

    # Card background & shadow outline
    c.setFillColor(colors.white)
    c.setStrokeColor(colors.HexColor("#E2E8F0"))
    c.setLineWidth(1.5)
    c.roundRect(qr_card_x, qr_card_y, qr_card_w, qr_card_h, 16, fill=1, stroke=1)

    # Card Top Badge
    c.setFillColor(colors.HexColor("#EFF6FF"))
    c.roundRect(qr_card_x + 30, qr_card_y + qr_card_h - 32, qr_card_w - 60, 24, 12, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#1D4ED8"))
    c.setFont("Helvetica-Bold", 10.5)
    c.drawCentredString(width / 2, qr_card_y + qr_card_h - 24, "POINT YOUR PHONE CAMERA TO SCAN")

    # Draw Big QR Code
    qr_size = 175
    qr_x = (width - qr_size) / 2
    qr_y = qr_card_y + 98
    c.drawImage(QR_PATH, qr_x, qr_y, width=qr_size, height=qr_size)

    # "Or Visit Our Website" box inside card
    c.setFillColor(colors.HexColor("#F8FAFC"))
    c.setStrokeColor(colors.HexColor("#CBD5E1"))
    c.setLineWidth(1)
    c.roundRect(qr_card_x + 18, qr_card_y + 14, qr_card_w - 36, 70, 10, fill=1, stroke=1)

    c.setFillColor(colors.HexColor("#64748B"))
    c.setFont("Helvetica", 9.5)
    c.drawCentredString(width / 2, qr_card_y + 64, "Prefer not to scan? Type this into your phone's browser:")

    c.setFillColor(colors.HexColor("#E4002B"))
    c.setFont("Helvetica-Bold", 18)
    c.drawCentredString(width / 2, qr_card_y + 40, "www.t3store.store")

    c.setFillColor(colors.HexColor("#0C1E3C"))
    c.setFont("Helvetica", 9)
    c.drawCentredString(width / 2, qr_card_y + 22, "Instant Booking  ·  View Real-Time Prices  ·  WhatsApp Order")

    # 4 Feature Bullet Boxes (2 columns x 2 rows)
    feat_y = qr_card_y - 20
    box_w = 265
    box_h = 42
    col1_x = margin + 20
    col2_x = width - margin - 20 - box_w

    features = [
        ("Crates & Packs Only", "Distributor wholesale prices for bulk orders", col1_x, feat_y - 42),
        ("Live Price Catalog", "Always see accurate, up-to-the-minute prices", col2_x, feat_y - 42),
        ("Fast Abuja Delivery", "Direct to your shop, home, party, or event", col1_x, feat_y - 92),
        ("Easy Online Booking", "Pick your drinks & confirm on WhatsApp instantly", col2_x, feat_y - 92),
    ]

    for title, desc, bx, by in features:
        c.setFillColor(colors.white)
        c.setStrokeColor(colors.HexColor("#E2E8F0"))
        c.setLineWidth(1)
        c.roundRect(bx, by, box_w, box_h, 8, fill=1, stroke=1)

        # Little red pill indicator
        c.setFillColor(colors.HexColor("#E4002B"))
        c.circle(bx + 16, by + box_h / 2, 4.5, fill=1, stroke=0)

        # Title
        c.setFillColor(colors.HexColor("#0C1E3C"))
        c.setFont("Helvetica-Bold", 10.5)
        c.drawString(bx + 28, by + 23, title)

        # Desc
        c.setFillColor(colors.HexColor("#64748B"))
        c.setFont("Helvetica", 8)
        c.drawString(bx + 28, by + 10, desc)

    # Footer Navy Contact Box
    footer_h = 78
    footer_y = margin + 5
    c.setFillColor(colors.HexColor("#0C1E3C"))
    c.rect(margin + 5, footer_y, width - 2 * (margin + 5), footer_h, fill=1, stroke=0)

    # Red accent line on footer top
    c.setFillColor(colors.HexColor("#E4002B"))
    c.rect(margin + 5, footer_y + footer_h - 4, width - 2 * (margin + 5), 4, fill=1, stroke=0)

    c.setFillColor(colors.HexColor("#F59E0B"))
    c.setFont("Helvetica-Bold", 12)
    c.drawCentredString(width / 2, footer_y + 50, "CALL / WHATSAPP:  0806 381 7772  ·  0708 627 7334")

    c.setFillColor(colors.white)
    c.setFont("Helvetica", 9)
    c.drawCentredString(width / 2, footer_y + 32, "Timber Market by Leisure Court Estate, Behind Aco Estate, Sabon Lugbe, Airport Road, Abuja")

    c.setFillColor(colors.HexColor("#94A3B8"))
    c.setFont("Helvetica", 8.5)
    c.drawCentredString(width / 2, footer_y + 16, "Opening Hours: Monday – Saturday  ·  8:00 AM – 8:00 PM")

    c.save()
    print(f"Generated PDF: {PDF_PATH}")

# 3. Generate Editable Word Document (.docx)
def generate_docx():
    doc = docx.Document()
    
    # Set narrow margins (0.4 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)

    # Header Table
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.5)
    
    # Navy background for header cell
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(r'<w:shd {} w:fill="0C1E3C"/>'.format(nsdecls('w')))
    tcPr.append(shd)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("T3 SUPERSTORE")
    run.font.name = "Arial"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(2)
    run2 = p2.add_run("WHOLESALE DRINKS, PACKS & CRATES")
    run2.font.name = "Arial"
    run2.font.size = Pt(13)
    run2.font.bold = True
    run2.font.color.rgb = RGBColor(245, 158, 11)  # Gold

    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(8)
    run3 = p3.add_run("Direct Distributor Prices  ·  Events & Wholesale  ·  Fast Abuja Delivery")
    run3.font.name = "Arial"
    run3.font.size = Pt(10)
    run3.font.color.rgb = RGBColor(226, 232, 240)

    # Announcement
    p_ann = doc.add_paragraph()
    p_ann.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ann.paragraph_format.space_before = Pt(14)
    p_ann.paragraph_format.space_after = Pt(2)
    run_ann = p_ann.add_run("★ YOU CAN NOW BOOK ONLINE! ★")
    run_ann.font.name = "Arial"
    run_ann.font.size = Pt(20)
    run_ann.font.bold = True
    run_ann.font.color.rgb = RGBColor(228, 0, 43)  # T3 Red

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    run_sub = p_sub.add_run("Browse our complete wholesale catalog with live prices and place orders directly from your phone.")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # QR Code Frame Table
    qr_tbl = doc.add_table(rows=1, cols=1)
    qr_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    qr_cell = qr_tbl.cell(0, 0)
    qr_cell.width = Inches(5.0)

    qr_tcPr = qr_cell._tc.get_or_add_tcPr()
    qr_shd = parse_xml(r'<w:shd {} w:fill="FFFFFF"/>'.format(nsdecls('w')))
    qr_tcPr.append(qr_shd)
    
    # Border
    borders = parse_xml(r'<w:tcBorders {}><w:top w:val="single" w:sz="12" w:space="0" w:color="CBD5E1"/><w:bottom w:val="single" w:sz="12" w:space="0" w:color="CBD5E1"/><w:left w:val="single" w:sz="12" w:space="0" w:color="CBD5E1"/><w:right w:val="single" w:sz="12" w:space="0" w:color="CBD5E1"/></w:tcBorders>'.format(nsdecls('w')))
    qr_tcPr.append(borders)

    p_scan = qr_cell.paragraphs[0]
    p_scan.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_scan.paragraph_format.space_before = Pt(8)
    p_scan.paragraph_format.space_after = Pt(6)
    r_scan = p_scan.add_run("📱 POINT YOUR PHONE CAMERA TO SCAN:")
    r_scan.font.name = "Arial"
    r_scan.font.size = Pt(12)
    r_scan.font.bold = True
    r_scan.font.color.rgb = RGBColor(12, 30, 60)

    p_img = qr_cell.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_after = Pt(8)
    p_img.add_run().add_picture(QR_PATH, width=Inches(2.5))

    p_link_lbl = qr_cell.add_paragraph()
    p_link_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_link_lbl.paragraph_format.space_after = Pt(2)
    r_lbl = p_link_lbl.add_run("Or type this into your phone's browser if you prefer not to scan:")
    r_lbl.font.name = "Arial"
    r_lbl.font.size = Pt(10)
    r_lbl.font.color.rgb = RGBColor(100, 116, 139)

    p_url = qr_cell.add_paragraph()
    p_url.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_url.paragraph_format.space_after = Pt(10)
    r_url = p_url.add_run("www.t3store.store")
    r_url.font.name = "Arial"
    r_url.font.size = Pt(18)
    r_url.font.bold = True
    r_url.font.color.rgb = RGBColor(228, 0, 43)

    # Feature List
    p_feat_title = doc.add_paragraph()
    p_feat_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_feat_title.paragraph_format.space_before = Pt(14)
    p_feat_title.paragraph_format.space_after = Pt(4)
    r_ft = p_feat_title.add_run("WHY ORDER WITH T3 SUPERSTORE ONLINE?")
    r_ft.font.name = "Arial"
    r_ft.font.size = Pt(11)
    r_ft.font.bold = True
    r_ft.font.color.rgb = RGBColor(12, 30, 60)

    features = [
        ("✔ Strictly Crates & Packs Only", "Direct distributor pricing on Coca-Cola, Pepsi, Maltina, Fearless, CWAY & more"),
        ("✔ Real-Time Wholesale Prices", "View updated prices anytime on your phone"),
        ("✔ Fast Delivery in Abuja", "Prompt supply to homes, shops, offices, weddings and events"),
        ("✔ Instant WhatsApp Order Confirmation", "Pick your products online and confirm your order in one click"),
    ]

    for f_title, f_desc in features:
        pf = doc.add_paragraph()
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.paragraph_format.left_indent = Inches(1.0)
        pf.paragraph_format.space_after = Pt(2)
        rt = pf.add_run(f_title + "  —  ")
        rt.font.name = "Arial"
        rt.font.size = Pt(10)
        rt.font.bold = True
        rt.font.color.rgb = RGBColor(12, 30, 60)

        rd = pf.add_run(f_desc)
        rd.font.name = "Arial"
        rd.font.size = Pt(9.5)
        rd.font.color.rgb = RGBColor(71, 85, 105)

    # Footer Table
    doc.add_paragraph().paragraph_format.space_before = Pt(8)
    ft_tbl = doc.add_table(rows=1, cols=1)
    ft_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_cell = ft_tbl.cell(0, 0)
    f_cell.width = Inches(7.5)
    f_tcPr = f_cell._tc.get_or_add_tcPr()
    f_shd = parse_xml(r'<w:shd {} w:fill="0C1E3C"/>'.format(nsdecls('w')))
    f_tcPr.append(f_shd)

    fp1 = f_cell.paragraphs[0]
    fp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp1.paragraph_format.space_before = Pt(8)
    fp1.paragraph_format.space_after = Pt(2)
    fr1 = fp1.add_run("📞 CALL / WHATSAPP:  0806 381 7772  ·  0708 627 7334")
    fr1.font.name = "Arial"
    fr1.font.size = Pt(12)
    fr1.font.bold = True
    fr1.font.color.rgb = RGBColor(245, 158, 11)

    fp2 = f_cell.add_paragraph()
    fp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp2.paragraph_format.space_after = Pt(2)
    fr2 = fp2.add_run("📍 Timber Market by Leisure Court Estate, Behind Aco Estate, Sabon Lugbe, Airport Road, Abuja")
    fr2.font.name = "Arial"
    fr2.font.size = Pt(9.5)
    fr2.font.color.rgb = RGBColor(255, 255, 255)

    fp3 = f_cell.add_paragraph()
    fp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp3.paragraph_format.space_after = Pt(8)
    fr3 = fp3.add_run("Opening Hours: Monday – Saturday  ·  8:00 AM – 8:00 PM")
    fr3.font.name = "Arial"
    fr3.font.size = Pt(9)
    fr3.font.color.rgb = RGBColor(148, 163, 184)

    doc.save(DOCX_PATH)
    print(f"Generated DOCX: {DOCX_PATH}")

# 4. Generate Web-Printable HTML Flyer
def generate_html():
    html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>T3 Superstore - Printable Poster / Flyer</title>
  <style>
    @page {{
      size: A4 portrait;
      margin: 10mm;
    }}
    * {{
      box-sizing: border-box;
      margin: 0;
      padding: 0;
      font-family: 'Segoe UI', -apple-system, BlinkMacSystemFont, Roboto, sans-serif;
    }}
    body {{
      background: #E2E8F0;
      display: flex;
      justify-content: center;
      padding: 20px;
    }}
    .sheet {{
      width: 210mm;
      min-height: 297mm;
      background: #FFFFFF;
      box-shadow: 0 10px 30px rgba(0,0,0,0.15);
      padding: 8mm;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
    }}
    .inner-border {{
      border: 2px solid #CBD5E1;
      border-radius: 12px;
      padding: 6mm;
      height: 100%;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
    }}
    .header {{
      background: #0C1E3C;
      color: white;
      text-align: center;
      padding: 18px 12px;
      border-radius: 10px;
      border-top: 5px solid #E4002B;
    }}
    .header h1 {{
      font-size: 32px;
      letter-spacing: 1px;
      font-weight: 900;
    }}
    .header .subtitle {{
      color: #F59E0B;
      font-weight: 700;
      font-size: 13px;
      letter-spacing: 2px;
      margin-top: 3px;
    }}
    .header .badges {{
      color: #E2E8F0;
      font-size: 11px;
      margin-top: 5px;
    }}
    .announcement {{
      text-align: center;
      margin-top: 14px;
    }}
    .pill {{
      display: inline-block;
      background: #E4002B;
      color: white;
      font-weight: 800;
      font-size: 11px;
      padding: 5px 18px;
      border-radius: 20px;
      letter-spacing: 1px;
    }}
    .announcement h2 {{
      font-size: 26px;
      color: #0C1E3C;
      margin-top: 8px;
      font-weight: 900;
    }}
    .announcement p {{
      color: #475569;
      font-size: 13px;
      margin-top: 4px;
      max-width: 520px;
      margin-left: auto;
      margin-right: auto;
    }}
    .qr-card {{
      background: #FFFFFF;
      border: 2px solid #E2E8F0;
      border-radius: 16px;
      padding: 16px;
      text-align: center;
      margin: 14px auto;
      max-width: 420px;
      box-shadow: 0 4px 15px rgba(0,0,0,0.04);
    }}
    .qr-badge {{
      background: #EFF6FF;
      color: #1D4ED8;
      font-weight: 800;
      font-size: 11.5px;
      padding: 4px 14px;
      border-radius: 12px;
      display: inline-block;
    }}
    .qr-img {{
      width: 200px;
      height: 200px;
      margin: 12px auto;
      display: block;
    }}
    .link-box {{
      background: #F8FAFC;
      border: 1px dashed #CBD5E1;
      border-radius: 10px;
      padding: 10px;
      margin-top: 8px;
    }}
    .link-box .small-note {{
      font-size: 11px;
      color: #64748B;
    }}
    .link-box .url {{
      font-size: 22px;
      font-weight: 900;
      color: #E4002B;
      margin: 3px 0;
      letter-spacing: 0.5px;
    }}
    .features-grid {{
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 8px;
      margin: 10px 0;
    }}
    .feature-item {{
      background: #F8FAFC;
      border: 1px solid #E2E8F0;
      border-radius: 8px;
      padding: 8px 12px;
      display: flex;
      align-items: center;
      gap: 8px;
    }}
    .dot {{
      width: 8px;
      height: 8px;
      background: #E4002B;
      border-radius: 50%;
      flex-shrink: 0;
    }}
    .feature-item h4 {{
      font-size: 11.5px;
      color: #0C1E3C;
      font-weight: 700;
    }}
    .feature-item p {{
      font-size: 9.5px;
      color: #64748B;
    }}
    .footer {{
      background: #0C1E3C;
      color: white;
      text-align: center;
      padding: 12px;
      border-radius: 10px;
      border-top: 4px solid #E4002B;
    }}
    .footer .phone {{
      color: #F59E0B;
      font-size: 14px;
      font-weight: 800;
    }}
    .footer .address {{
      font-size: 10.5px;
      margin-top: 4px;
      color: #F1F5F9;
    }}
    .footer .hours {{
      font-size: 9.5px;
      color: #94A3B8;
      margin-top: 3px;
    }}
    .print-btn {{
      position: fixed;
      top: 20px;
      right: 20px;
      background: #E4002B;
      color: white;
      border: none;
      padding: 12px 24px;
      font-size: 14px;
      font-weight: bold;
      border-radius: 30px;
      cursor: pointer;
      box-shadow: 0 4px 15px rgba(228,0,43,0.4);
    }}
    @media print {{
      body {{
        background: none;
        padding: 0;
      }}
      .sheet {{
        box-shadow: none;
        width: 100%;
        min-height: 100%;
        padding: 0;
      }}
      .print-btn {{
        display: none;
      }}
    }}
  </style>
</head>
<body>
  <button class="print-btn" onclick="window.print()">🖨️ Print Poster / Flyer</button>

  <div class="sheet">
    <div class="inner-border">
      <!-- Header Banner -->
      <div class="header">
        <h1>T3 SUPERSTORE</h1>
        <div class="subtitle">WHOLESALE DRINKS, PACKS & CRATES</div>
        <div class="badges">Direct Distributor Prices &middot; Events & Wholesale &middot; Fast Abuja Delivery</div>
      </div>

      <!-- Main Callout -->
      <div class="announcement">
        <div class="pill">&starf; OFFICIAL ANNOUNCEMENT &starf;</div>
        <h2>YOU CAN NOW BOOK ONLINE!</h2>
        <p>Browse our complete catalog with live prices and place orders directly from your smartphone.</p>
      </div>

      <!-- Big QR Code Container -->
      <div class="qr-card">
        <div class="qr-badge">&camera; POINT YOUR PHONE CAMERA TO SCAN</div>
        <img src="{QR_PATH}" alt="Scan to visit www.t3store.store" class="qr-img">
        <div class="link-box">
          <div class="small-note">Prefer not to scan? Type this into your phone's browser:</div>
          <div class="url">www.t3store.store</div>
          <div class="small-note">Instant Online Ordering &middot; Live Prices &middot; WhatsApp Support</div>
        </div>
      </div>

      <!-- Features Grid -->
      <div class="features-grid">
        <div class="feature-item">
          <div class="dot"></div>
          <div>
            <h4>Strictly Crates & Packs Only</h4>
            <p>Direct wholesale prices on full crates and packs</p>
          </div>
        </div>
        <div class="feature-item">
          <div class="dot"></div>
          <div>
            <h4>Live Price Catalog</h4>
            <p>Always view updated wholesale prices in real-time</p>
          </div>
        </div>
        <div class="feature-item">
          <div class="dot"></div>
          <div>
            <h4>Fast Delivery Across Abuja</h4>
            <p>Direct supply to homes, shops, offices and events</p>
          </div>
        </div>
        <div class="feature-item">
          <div class="dot"></div>
          <div>
            <h4>Easy 1-Click WhatsApp Booking</h4>
            <p>Pick your items online and confirm your order right away</p>
          </div>
        </div>
      </div>

      <!-- Footer Contacts -->
      <div class="footer">
        <div class="phone">&phone; CALL / WHATSAPP: 0806 381 7772 &middot; 0708 627 7334</div>
        <div class="address">&map-pin; Timber Market by Leisure Court Estate, Behind Aco Estate, Sabon Lugbe, Airport Road, Abuja</div>
        <div class="hours">&clock; Monday &ndash; Saturday &middot; 8:00 AM &ndash; 8:00 PM</div>
      </div>
    </div>
  </div>
</body>
</html>
"""
    with open(HTML_PATH, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Generated HTML: {HTML_PATH}")

if __name__ == "__main__":
    generate_qr()
    generate_pdf()
    generate_docx()
    generate_html()
    print("ALL PRINTABLE ASSETS GENERATED SUCCESSFULLY!")
